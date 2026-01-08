import json
import os
import pandas as pd
from jinja2 import Environment, FileSystemLoader
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

class MarkdownGenerator:
    def __init__(self, template_dir):
        self.env = Environment(loader=FileSystemLoader(template_dir))

    def generate(self, data, template_name, output_path):
        template = self.env.get_template(template_name)
        rendered_content = template.render(data)
        
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(rendered_content)
        return output_path

class CSVGenerator:
    def generate(self, data, output_dir):
        os.makedirs(output_dir, exist_ok=True)
        paths = []
        
        # Example: Success Rate table
        if 'rules' in data and 'success_rate' in data['rules']:
            df = pd.DataFrame(data['rules']['success_rate'])
            path = os.path.join(output_dir, 'success_rate.csv')
            df.to_csv(path, index=False, encoding='utf-8-sig')
            paths.append(path)
            
        return paths

class ExcelGenerator:
    def generate(self, data, output_path):
        wb = Workbook()
        ws = wb.active
        ws.title = "Summary"

        # Styles
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
        
        # 1. Summary Sheet
        ws['A1'] = "Category"
        ws['B1'] = "Value"
        for cell in ['A1', 'B1']:
            ws[cell].font = header_font
            ws[cell].fill = header_fill

        ws.append(["Title", data.get("title", "")])
        ws.append(["Summary", data.get("summary", "")])
        ws.append(["Max Level", data.get("rules", {}).get("max_level", "")])

        # 2. Success Rates Sheet
        if 'rules' in data and 'success_rate' in data['rules']:
            ws_rates = wb.create_sheet("Success Rates")
            headers = ["Level", "Rate (%)"]
            ws_rates.append(headers)
            for cell in ws_rates[1]:
                cell.font = header_font
                cell.fill = header_fill

            for item in data['rules']['success_rate']:
                ws_rates.append([item['level'], item['rate'] * 100])

        # Auto-adjust column width
        for sheet in wb.worksheets:
            for col in sheet.columns:
                max_length = 0
                column = col[0].column_letter
                for cell in col:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = (max_length + 2)
                sheet.column_dimensions[column].width = adjusted_width

        wb.save(output_path)
        return output_path

from docx import Document
from docx.shared import Inches, Pt

class DocxGenerator:
    def generate(self, data, output_path):
        doc = Document()
        
        # Title
        doc.add_heading(data.get('title', 'Design Document'), 0)
        
        # Summary
        doc.add_heading('Summary', level=1)
        doc.add_paragraph(data.get('summary', ''))
        
        # Rules / Tables
        if 'rules' in data:
            doc.add_heading('System Rules', level=1)
            rules = data['rules']
            
            # Max Level
            p = doc.add_paragraph()
            p.add_run('Max Level: ').bold = True
            p.add_run(str(rules.get('max_level', 'N/A')))
            
            # Success Rate Table
            if 'success_rate' in rules:
                doc.add_heading('Success Rate Table', level=2)
                rates = rules['success_rate']
                table = doc.add_table(rows=1, cols=2)
                table.style = 'Table Grid'
                hdr_cells = table.rows[0].cells
                hdr_cells[0].text = 'Level'
                hdr_cells[1].text = 'Success Rate'
                
                for item in rates:
                    row_cells = table.add_row().cells
                    row_cells[0].text = str(item.get('level', ''))
                    row_cells[1].text = str(item.get('rate', ''))
                    
        # Costs
        if 'costs' in data:
            doc.add_heading('Costs', level=1)
            costs = data['costs']
            for key, val in costs.items():
                p = doc.add_paragraph(style='List Bullet')
                p.add_run(f"{key}: ").bold = True
                p.add_run(str(val))
                
        # Exceptions
        if 'exceptions' in data:
            doc.add_heading('Exceptions', level=1)
            for exc in data['exceptions']:
                doc.add_paragraph(exc, style='List Bullet')
                
        doc.save(output_path)
        return output_path
